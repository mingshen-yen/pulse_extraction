#!/usr/bin/env python3
"""pulse_table.xlsx  ->  D1 workbook tables (one table per sheet).

D1 is the source of truth for the reference catalogs; this script is for the
initial load and for replacing the tables wholesale from a new workbook.  It
DROPS and recreates every workbook table, so any edits made in D1 since the
last import are lost -- export a backup first:

    wrangler d1 export pulse_api --remote --output backup.sql
    python scripts/import_workbook.py path/to/pulse_table.xlsx      # -> db/workbook.sql
    wrangler d1 execute pulse_api --remote --file db/workbook.sql --yes

Column names are the sheet's own header cells, verbatim (quote them in SQL:
"Tp (s)").  Every table gets `_row`, the 1-based data-row number in the sheet,
which keeps the workbook's order.  Booleans become 1/0, dates ISO 8601 text.
"""

from __future__ import annotations

import argparse
import datetime as dt
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "db" / "workbook.sql"

# sheet name -> D1 table name
TABLES = {
    "Overview": "overview",
    "Sources_method": "sources_method",
    "Station_coords": "station_coords",
    "Event_sources": "event_sources",
    "Finite_fault_segments": "finite_fault_segments",
    "Event_match_audit": "event_match_audit",
    "Pulse_records": "pulse_records",
    "S&B": "sheet_sb",
    "NCREE": "sheet_ncree",
    "YEN": "sheet_yen",
    "Distance_check": "distance_check",
}
# lookups the API makes (functions/api/_lib.js)
INDEXES = {
    "pulse_records": ["active_sheet", "event_source_key"],
    "event_sources": ["event_source_key"],
    "finite_fault_segments": ["event_key"],
}


def ident(name: str) -> str:
    return '"' + name.replace('"', '""') + '"'


def value(v):
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, (dt.datetime, dt.date, dt.time)):
        return v.isoformat()
    return v


def literal(v) -> str:
    if v is None:
        return "NULL"
    if isinstance(v, (int, float)):
        return repr(v)
    return "'" + str(v).replace("'", "''") + "'"


def affinity(values) -> str:
    """INTEGER / REAL / TEXT when a column is uniform, else no declared type."""
    kinds = {type(v) for v in values if v is not None}
    if not kinds:
        return ""
    if kinds <= {int}:
        return " INTEGER"
    if kinds <= {int, float}:
        return " REAL"
    if kinds <= {str}:
        return " TEXT"
    return ""


def read_sheet(ws):
    rows = list(ws.iter_rows(values_only=True))
    # rows can be ragged (trailing empty cells are left out); square them up
    width = max(len(r) for r in rows)
    rows = [tuple(r) + (None,) * (width - len(r)) for r in rows]
    head = [str(h).strip() if h not in (None, "") else "" for h in rows[0]]
    # blank header cells (the Overview sheet) get positional names; SQLite
    # column names are case-insensitive, so a repeat such as the S&B sheet's
    # "Mechanism" / "mechanism" gets its column number appended
    seen, cols = set(), []
    for i, h in enumerate(head, start=1):
        name = f"col{i}" if not h else h if h.lower() not in seen else f"{h}_col{i}"
        seen.add(name.lower())
        cols.append(name)
    data = [[value(v) for v in r] for r in rows[1:]]
    # drop trailing all-empty rows
    while data and all(v is None for v in data[-1]):
        data.pop()
    return cols, data


def sheet_sql(table: str, cols: list[str], data: list[list]) -> list[str]:
    defs = ", ".join(f"{ident(c)}{affinity(r[i] for r in data)}" for i, c in enumerate(cols))
    out = [f"DROP TABLE IF EXISTS {ident(table)};",
           f"CREATE TABLE {ident(table)} (_row INTEGER PRIMARY KEY, {defs});"]
    names = ", ".join(["_row", *map(ident, cols)])
    for n, r in enumerate(data, start=1):
        out.append(f"INSERT INTO {ident(table)} ({names}) VALUES "
                   f"({n}, {', '.join(literal(v) for v in r)});")
    for col in INDEXES.get(table, []):
        out.append(f"CREATE INDEX {ident(f'{table}_{col}')} ON {ident(table)} ({ident(col)});")
    return out


def build(xlsx: Path) -> tuple[list[str], dict[str, int]]:
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    missing = set(TABLES) - set(wb.sheetnames)
    if missing:
        raise SystemExit(f"{xlsx}: missing sheet(s) {sorted(missing)}")
    stmts, counts = [], {}
    for sheet, table in TABLES.items():
        cols, data = read_sheet(wb[sheet])
        stmts += sheet_sql(table, cols, data)
        counts[table] = len(data)
    return stmts, counts


def main(argv=None) -> None:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx", type=Path)
    ap.add_argument("-o", "--out", type=Path, default=OUT)
    args = ap.parse_args(argv)

    stmts, counts = build(args.xlsx)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text("\n".join(stmts) + "\n")
    for table, n in counts.items():
        print(f"{table:<24} {n:5d} rows")
    print(f"wrote {args.out}  ({len(stmts)} statements)")


if __name__ == "__main__":
    main()
