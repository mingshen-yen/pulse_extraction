"""scripts/import_workbook.py: a workbook becomes one SQLite/D1 table per sheet."""

import sqlite3
import sys
from datetime import date
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import import_workbook  # noqa: E402


@pytest.fixture
def workbook(tmp_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet in import_workbook.TABLES:
        wb.create_sheet(sheet).append(["id"])
    pr = wb["Pulse_records"]
    pr.delete_rows(1)
    pr.append(["record_id", "Station Name", "Ipulse_H", "Tp (s)", "note", "Mechanism",
               "mechanism"])
    pr.append([1, "O'Neill", True, 1.5, date(2023, 2, 6), "SS", "strike slip"])
    pr.append([2, "TCU052", False, None])          # ragged: trailing cells missing
    pr.append([None] * 7)                           # trailing empty row (formatting)
    wb["Overview"].append([None, "text in an untitled column"])
    path = tmp_path / "t.xlsx"
    wb.save(path)
    return path


def _load(path, tmp_path):
    out = tmp_path / "w.sql"
    import_workbook.main([str(path), "-o", str(out)])
    db = sqlite3.connect(":memory:")
    db.executescript(out.read_text())
    return db


def test_every_sheet_becomes_a_table(workbook, tmp_path):
    db = _load(workbook, tmp_path)
    tables = {r[0] for r in db.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    assert set(import_workbook.TABLES.values()) <= tables


def test_values_and_columns_survive(workbook, tmp_path):
    db = _load(workbook, tmp_path)
    db.row_factory = sqlite3.Row
    rows = db.execute("SELECT * FROM pulse_records ORDER BY _row").fetchall()
    assert len(rows) == 2                            # empty trailing row dropped
    a, b = rows
    assert (a["_row"], a["Station Name"], a["Ipulse_H"], a["Tp (s)"]) == (1, "O'Neill", 1, 1.5)
    assert a["note"] == "2023-02-06T00:00:00"
    assert (a["Mechanism"], a["mechanism_col7"]) == ("SS", "strike slip")
    assert (b["Ipulse_H"], b["Tp (s)"], b["note"]) == (0, None, None)
    assert db.execute("SELECT col2 FROM overview").fetchone()[0] == "text in an untitled column"


def test_missing_sheet_is_an_error(tmp_path):
    wb = openpyxl.Workbook()
    path = tmp_path / "bad.xlsx"
    wb.save(path)
    with pytest.raises(SystemExit):
        import_workbook.build(path)
